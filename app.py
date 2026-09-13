import os
import re
import io
import json
import html
import time
from collections import Counter
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

import numpy as np
import pandas as pd
import plotly.express as px
import requests
import streamlit as st
from dotenv import load_dotenv


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="ScholarReady AI",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .block-container {
        max-width: 1450px;
        padding-top: 1.5rem;
        padding-bottom: 3rem;
    }
    [data-testid="stMetric"] {
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        border-radius: 14px;
        padding: 14px;
    }
    [data-testid="stSidebar"] {
        background: #f8fafc;
    }
    h1, h2, h3 {
        letter-spacing: -0.02em;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# SETTINGS AND SECRETS
# =========================================================

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
CURRENT_DATE = date.today()

GROQ_PROFILE_MODEL_DEFAULT = "openai/gpt-oss-120b"
GROQ_JUDGE_MODEL_DEFAULT = "openai/gpt-oss-20b"
GEMINI_MODEL_DEFAULT = "gemini-flash-latest"
EMBEDDING_MODEL = "all-MiniLM-L6-v2"


def get_secret(name, default=""):
    try:
        value = st.secrets.get(name, "")
        if value:
            return str(value)
    except Exception:
        pass
    return os.environ.get(name, default)


GROQ_API_KEY = get_secret("GROQ_API_KEY")
TAVILY_API_KEY = get_secret("TAVILY_API_KEY")
FIRECRAWL_API_KEY = get_secret("FIRECRAWL_API_KEY")
GEMINI_API_KEY = get_secret("GEMINI_API_KEY")

GROQ_PROFILE_MODEL = get_secret(
    "GROQ_PROFILE_MODEL", GROQ_PROFILE_MODEL_DEFAULT
)
GROQ_JUDGE_MODEL = get_secret(
    "GROQ_JUDGE_MODEL", GROQ_JUDGE_MODEL_DEFAULT
)
GEMINI_MODEL = get_secret(
    "GEMINI_MODEL", GEMINI_MODEL_DEFAULT
)


# =========================================================
# GENERAL UTILITIES
# =========================================================

def safe_float(value, default=0.0):
    try:
        if pd.isna(value):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def safe_int(value, default=0):
    try:
        if pd.isna(value):
            return default
        return int(float(value))
    except (TypeError, ValueError):
        return default


def parse_boolean(value):
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "1", "yes", "y"}


def normalize_column_name(value):
    value = str(value).strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")


def split_comma_values(text):
    return [
        value.strip()
        for value in str(text or "").split(",")
        if value.strip()
    ]


def split_line_values(text):
    return [
        value.strip()
        for value in str(text or "").splitlines()
        if value.strip()
    ]


def find_csv_file():
    preferred = [
        BASE_DIR / "scholarships_final.csv",
        BASE_DIR / "scholarships_merged_full.csv",
    ]
    for path in preferred:
        if path.exists():
            return path
    available = list(BASE_DIR.glob("*.csv"))
    if available:
        return max(available, key=lambda path: path.stat().st_size)
    raise FileNotFoundError(
        "No scholarship CSV was found beside app.py. "
        "Upload scholarships_final.csv to GitHub."
    )


# =========================================================
# DEGREE-LEVEL DETECTION
# =========================================================

def detect_scholarship_levels(years_value):
    text = html.unescape(str(years_value or "")).lower()
    levels = set()

    if re.search(r"\bpostdoctoral\b|\bpostdoc\b", text):
        levels.add("POSTDOC")

    if re.search(
        r"\bdoctoral[- ]level\b|\bdoctoral\b|\bdoctorate\b|\bph\.?d\b",
        text,
    ):
        levels.add("PhD")

    if re.search(
        r"\bmaster'?s[- ]level\b|\bmaster'?s degree\b|"
        r"\bmaster'?s study\b|\bmasters?\b|\bmsc\b|\bm\.sc\b|\bmba\b",
        text,
    ):
        levels.add("MS")

    if re.search(
        r"\bhigh school senior\b|\bcollege freshman\b|"
        r"\bcollege sophomore\b|\bcollege junior\b|"
        r"\bcollege senior\b|\bundergraduate\b|\bbachelor'?s\b",
        text,
    ):
        levels.add("BS")

    if not levels:
        levels.add("UNCLEAR")

    ordered = ["BS", "MS", "PhD", "POSTDOC", "UNCLEAR"]
    return tuple(level for level in ordered if level in levels)


def level_is_allowed(levels, selected_level, include_unclear=False):
    levels = set(levels or [])
    if include_unclear and "UNCLEAR" in levels:
        return True
    if selected_level == "Any":
        return bool(levels.intersection({"BS", "MS", "PhD"}))
    return selected_level in levels


# =========================================================
# STRATEGIC RECORDS
# =========================================================

def strategic_anchor_records():
    return [
        {
            "scholarship_name":
                "KAUST Fellowship for Admitted MS and PhD Students",
            "provider_name":
                "King Abdullah University of Science and Technology",
            "provider_type": "University",
            "deadline": "LIVE VERIFY CURRENT ADMISSION ROUND",
            "amount": 0,
            "funding_summary":
                "Graduate fellowship consideration; verify current benefits",
            "description":
                "International graduate opportunity at KAUST in "
                "artificial intelligence, computer science, machine "
                "learning, statistics, applied mathematics, electrical "
                "engineering, computer vision, data science and related "
                "STEM research fields.",
            "location": "Saudi Arabia; international applicants",
            "country": "Saudi Arabia",
            "years": "Master's-level study|Doctoral-level study",
            "field_tags":
                "Artificial Intelligence|Computer Science|Machine Learning|"
                "Data Science|Computer Vision|Engineering|Research|STEM",
            "link": "https://admissions.kaust.edu.sa/study",
            "source": "strategic_anchor",
            "verified": False,
            "official_seed": True,
            "is_gold": True,
            "needs_verify": True,
        },
        {
            "scholarship_name":
                "Erasmus Mundus Joint Masters Programme Scholarships",
            "provider_name":
                "European Commission and participating universities",
            "provider_type": "University consortium",
            "deadline": "VARIES BY INDIVIDUAL PROGRAMME — LIVE VERIFY",
            "amount": 0,
            "funding_summary": "Competitive programme-specific scholarships",
            "description":
                "International joint master's programmes delivered by "
                "university consortia. Applicants must choose a specific "
                "programme because entry requirements, language tests, "
                "documents, deadlines and funding vary. Relevant "
                "programmes may exist in artificial intelligence, "
                "computer science, data science, computer vision, NLP "
                "and robotics.",
            "location": "Multiple European and partner countries",
            "country": "Multiple European Countries",
            "years": "Master's-level study",
            "field_tags":
                "Artificial Intelligence|Computer Science|Machine Learning|"
                "Data Science|Computer Vision|NLP|Robotics|Engineering",
            "link":
                "https://erasmus-plus.ec.europa.eu/opportunities/"
                "individuals/students/erasmus-mundus-joint-masters",
            "source": "strategic_anchor",
            "verified": False,
            "official_seed": True,
            "is_gold": True,
            "needs_verify": True,
        },
    ]


# =========================================================
# DATA LOADING
# =========================================================

@st.cache_data(show_spinner=False)
def load_data():
    csv_path = find_csv_file()
    dataframe = pd.read_csv(csv_path, low_memory=False)

    dataframe.columns = [
        normalize_column_name(column) for column in dataframe.columns
    ]

    aliases = {
        "name": "scholarship_name",
        "title": "scholarship_name",
        "scholarship": "scholarship_name",
        "url": "link",
        "website": "link",
        "website_link": "link",
        "degree": "years",
        "degree_level": "years",
        "degree_levels": "years",
        "award_amount": "amount",
    }

    for old_name, new_name in aliases.items():
        if new_name not in dataframe.columns and old_name in dataframe.columns:
            dataframe[new_name] = dataframe[old_name]

    defaults = {
        "scholarship_name": "",
        "provider_name": "",
        "provider_type": "Unknown",
        "deadline": "Unknown deadline",
        "amount": 0,
        "funding_summary": "",
        "description": "",
        "location": "",
        "country": "",
        "years": "",
        "field_tags": "",
        "link": "",
        "source": "raw_csv",
        "verified": False,
        "official_seed": False,
        "is_gold": False,
        "needs_verify": False,
    }

    for column, default in defaults.items():
        if column not in dataframe.columns:
            dataframe[column] = default
        else:
            dataframe[column] = dataframe[column].fillna(default)

    for column in [
        "scholarship_name", "provider_name", "provider_type",
        "deadline", "funding_summary", "description", "location",
        "country", "years", "field_tags", "link", "source",
    ]:
        dataframe[column] = (
            dataframe[column]
            .fillna("")
            .astype(str)
            .apply(html.unescape)
            .str.strip()
        )

    dataframe["amount"] = pd.to_numeric(
        dataframe["amount"], errors="coerce"
    ).fillna(0)

    for column in ["verified", "official_seed", "is_gold", "needs_verify"]:
        dataframe[column] = dataframe[column].apply(parse_boolean)

    normalized_names = (
        dataframe["scholarship_name"]
        .str.lower()
        .str.replace(r"[^a-z0-9]+", " ", regex=True)
        .str.strip()
    )

    generic_names = {
        "kaust fellowship",
        "kaust fellowship for admitted ms and phd students",
        "erasmus mundus joint masters",
        "erasmus mundus joint master degrees",
        "erasmus mundus joint masters programme scholarships",
    }

    dataframe = dataframe.loc[~normalized_names.isin(generic_names)].copy()

    anchors = pd.DataFrame(strategic_anchor_records())

    all_columns = list(
        dict.fromkeys(list(anchors.columns) + list(dataframe.columns))
    )

    anchors = anchors.reindex(columns=all_columns)
    dataframe = dataframe.reindex(columns=all_columns)

    dataframe = pd.concat([anchors, dataframe], ignore_index=True)

    for column, default in defaults.items():
        if column not in dataframe.columns:
            dataframe[column] = default
        else:
            dataframe[column] = dataframe[column].fillna(default)

    for column in ["verified", "official_seed", "is_gold", "needs_verify"]:
        dataframe[column] = dataframe[column].apply(parse_boolean)

    source_text = dataframe["source"].fillna("").astype(str).str.lower()

    dataframe["is_curated"] = (
        dataframe["official_seed"]
        | dataframe["is_gold"]
        | source_text.str.contains(
            r"strategic_anchor|ai_enriched|curated|official|manual|gold",
            regex=True, na=False,
        )
    )

    unknown_deadline = (
        dataframe["deadline"]
        .astype(str)
        .str.contains(
            r"unknown|verify|varies|not available",
            case=False, regex=True, na=True,
        )
        | dataframe["deadline"].astype(str).str.strip().eq("")
    )

    suspicious_amount = dataframe["amount"].isin([0, 500])

    dataframe["data_quality_problem"] = (
        ~dataframe["official_seed"]
        & (unknown_deadline | suspicious_amount)
    )

    dataframe["needs_verify"] = (
        dataframe["needs_verify"]
        | ~dataframe["verified"]
        | unknown_deadline
        | suspicious_amount
    )

    dataframe["_levels"] = dataframe["years"].apply(detect_scholarship_levels)

    dataframe["search_text"] = (
        "Scholarship: " + dataframe["scholarship_name"].astype(str)
        + ". Provider: " + dataframe["provider_name"].astype(str)
        + ". Description: " + dataframe["description"].astype(str)
        + ". Fields: " + dataframe["field_tags"].astype(str)
        + ". Degree level: " + dataframe["years"].astype(str)
        + ". Location: " + dataframe["location"].astype(str)
    )

    dataframe = (
        dataframe
        .drop_duplicates(subset=["scholarship_name", "link"], keep="first")
        .reset_index(drop=True)
    )

    dataframe["record_id"] = [
        f"SCH-{index:06d}" for index in range(len(dataframe))
    ]

    return dataframe, str(csv_path)


try:
    DATA, DATA_FILE = load_data()
except Exception as error:
    st.error(f"Scholarship dataset could not be loaded: {error}")
    st.stop()


# =========================================================
# FIELD NORMALIZATION
# =========================================================

FIELD_OPTIONS = [
    "Artificial Intelligence", "Computer Science", "Machine Learning",
    "Data Science", "Computer Vision", "Natural Language Processing",
    "Robotics", "Engineering", "Cybersecurity", "Medicine",
    "Public Health", "Business", "Law", "Education", "Agriculture", "Arts",
]


def normalize_fields(values, major="", profession=""):
    if isinstance(values, str):
        values = re.split(r"[,/|;]+", values)

    combined = list(values or [])
    combined.extend([major, profession])

    mappings = [
        (r"\bartificial intelligence\b|\bai\b", "Artificial Intelligence"),
        (r"\bcomputer science\b|\bcs\b|software", "Computer Science"),
        (r"\bmachine learning\b|\bml\b|deep learning", "Machine Learning"),
        (r"\bdata science\b|data analysis|data analyst", "Data Science"),
        (r"\bcomputer vision\b|image processing", "Computer Vision"),
        (r"natural language processing|\bnlp\b", "Natural Language Processing"),
        (r"\brobotics?\b", "Robotics"),
        (r"\bcybersecurity\b|cyber security", "Cybersecurity"),
        (r"\bengineering\b", "Engineering"),
        (r"\bpublic health\b", "Public Health"),
        (r"\bmedicine\b|\bmedical\b", "Medicine"),
        (r"\bbusiness\b|\bmba\b|management", "Business"),
        (r"\blaw\b|\blegal\b", "Law"),
        (r"\beducation\b|\bteaching\b", "Education"),
        (r"\bagriculture\b|agronomy|horticulture", "Agriculture"),
        (r"\barts?\b|design|music|film", "Arts"),
    ]

    output = []
    for value in combined:
        value_text = str(value or "").replace("_", " ").strip().lower()
        for pattern, normalized in mappings:
            if re.search(pattern, value_text):
                output.append(normalized)

    return list(dict.fromkeys(output))


# =========================================================
# DIRECT CV FACT EXTRACTION
# =========================================================

def extract_gpa_from_cv_text(cv_text):
    cleaned = html.unescape(str(cv_text or ""))

    patterns = [
        r"\bCGPA\s*(?:is|of|:|=)?\s*([0-9]+(?:\.[0-9]+)?)"
        r"(?:\s*(?:/|out\s+of)\s*([0-9]+(?:\.[0-9]+)?))?",
        r"\bcumulative\s+(?:GPA|grade\s+point\s+average)\s*"
        r"(?:is|of|:|=)?\s*([0-9]+(?:\.[0-9]+)?)"
        r"(?:\s*(?:/|out\s+of)\s*([0-9]+(?:\.[0-9]+)?))?",
        r"\bGPA\s*(?:is|of|:|=)?\s*([0-9]+(?:\.[0-9]+)?)"
        r"(?:\s*(?:/|out\s+of)\s*([0-9]+(?:\.[0-9]+)?))?",
    ]

    for pattern in patterns:
        match = re.search(pattern, cleaned, re.IGNORECASE)
        if not match:
            continue

        value = safe_float(match.group(1), 0.0)
        scale = safe_float(match.group(2), 0.0)

        if value <= 0:
            continue

        if scale > 0:
            if value <= scale:
                normalized = value / scale * 4.0
                if 0 < normalized <= 4:
                    return round(normalized, 2)
        elif value <= 4:
            return round(value, 2)

    return 0.0


def extract_graduation_year(cv_text):
    years = [
        int(year)
        for year in re.findall(r"\b20(?:2[4-9]|3[0-5])\b", str(cv_text or ""))
    ]
    if not years:
        return 0
    future_years = [year for year in years if year >= CURRENT_DATE.year]
    return max(future_years) if future_years else max(years)


def fallback_name(cv_text):
    for line in str(cv_text).splitlines():
        line = line.strip()
        if not line:
            continue
        lower = line.lower()
        if any(
            marker in lower
            for marker in [
                "@", "linkedin", "objective", "resume",
                "curriculum", "phone", "profile",
            ]
        ):
            continue
        if 2 <= len(line.split()) <= 5 and len(line) <= 60 and re.search(
            r"[A-Za-z]", line
        ):
            return line
    return "Student"


def detect_current_degree(cv_text):
    text = str(cv_text).lower()
    if re.search(
        r"bachelor of science|bachelor'?s|\bbsai\b|\bbsc\b|\bb\.sc\b", text
    ):
        return "BS"
    if re.search(r"master of science|master'?s|\bmsc\b|\bm\.sc\b", text):
        return "MS"
    if re.search(r"\bphd\b|\bph\.d\b|\bdoctoral\b", text):
        return "PhD"
    return "unknown"


# =========================================================
# JSON SCHEMAS
# =========================================================

CV_SCHEMA = {
    "type": "object",
    "properties": {
        "name": {"type": "string"},
        "gender": {
            "type": "string",
            "enum": ["female", "male", "non-binary", "unknown"],
        },
        "nationality": {"type": "string"},
        "current_degree": {
            "type": "string",
            "enum": ["BS", "MS", "PhD", "unknown"],
        },
        "target_degree": {"type": "string", "enum": ["BS", "MS", "PhD"]},
        "expected_graduation_year": {"type": "integer"},
        "gpa": {"type": "number"},
        "major": {"type": "string"},
        "profession": {"type": "string"},
        "career_goal": {"type": "string"},
        "field_of_study": {"type": "array", "items": {"type": "string"}},
        "interests": {"type": "array", "items": {"type": "string"}},
        "skills": {"type": "array", "items": {"type": "string"}},
        "research_experience": {"type": "boolean"},
        "research_topics": {"type": "array", "items": {"type": "string"}},
        "projects": {"type": "array", "items": {"type": "string"}},
        "work_experience_years": {"type": "number"},
        "leadership": {"type": "array", "items": {"type": "string"}},
        "certifications": {"type": "array", "items": {"type": "string"}},
        "achievements": {"type": "array", "items": {"type": "string"}},
        "strengths": {"type": "array", "items": {"type": "string"}},
        "gaps": {"type": "array", "items": {"type": "string"}},
    },
    "required": [
        "name", "gender", "nationality", "current_degree", "target_degree",
        "expected_graduation_year", "gpa", "major", "profession",
        "career_goal", "field_of_study", "interests", "skills",
        "research_experience", "research_topics", "projects",
        "work_experience_years", "leadership", "certifications",
        "achievements", "strengths", "gaps",
    ],
    "additionalProperties": False,
}


JUDGE_SCHEMA = {
    "type": "object",
    "properties": {
        "results": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "candidate_id": {"type": "integer"},
                    "status": {
                        "type": "string",
                        "enum": [
                            "strong", "possible", "verify_first",
                            "wrong_field", "not_eligible", "suspicious",
                        ],
                    },
                    "fit_score": {"type": "integer"},
                    "why": {"type": "string"},
                    "missing": {"type": "array", "items": {"type": "string"}},
                    "actions": {"type": "array", "items": {"type": "string"}},
                    "warnings": {"type": "array", "items": {"type": "string"}},
                    "priority": {
                        "type": "string",
                        "enum": ["high", "medium", "low"],
                    },
                },
                "required": [
                    "candidate_id", "status", "fit_score", "why",
                    "missing", "actions", "warnings", "priority",
                ],
                "additionalProperties": False,
            },
        },
    },
    "required": ["results"],
    "additionalProperties": False,
}


VERIFY_SCHEMA = {
    "type": "object",
    "properties": {
        "scholarship_name": {"type": "string"},
        "current_deadline": {"type": "string"},
        "current_amount": {"type": "string"},
        "application_status": {
            "type": "string",
            "enum": ["open", "closed", "upcoming", "unknown"],
        },
        "offered_levels": {"type": "array", "items": {"type": "string"}},
        "eligible_fields": {"type": "array", "items": {"type": "string"}},
        "requirements": {"type": "array", "items": {"type": "string"}},
        "documents_needed": {"type": "array", "items": {"type": "string"}},
        "source_url": {"type": "string"},
        "source_type": {
            "type": "string",
            "enum": ["official", "aggregator", "unclear"],
        },
        "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
        "student_verdict": {
            "type": "string",
            "enum": ["ready", "almost", "not_eligible", "wrong_field", "unclear"],
        },
        "student_fit_score": {"type": "integer"},
        "why_for_student": {"type": "string"},
        "missing_for_student": {"type": "array", "items": {"type": "string"}},
    },
    "required": [
        "scholarship_name", "current_deadline", "current_amount",
        "application_status", "offered_levels", "eligible_fields",
        "requirements", "documents_needed", "source_url", "source_type",
        "confidence", "student_verdict", "student_fit_score",
        "why_for_student", "missing_for_student",
    ],
    "additionalProperties": False,
}


# =========================================================
# AI CALLS
# =========================================================

def extract_json(text):
    text = str(text or "").strip()
    if "```" in text:
        blocks = text.split("```")
        if len(blocks) > 1:
            text = blocks[1]
        if text.startswith("json"):
            text = text[4:]
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        text = text[start:end + 1]
    return json.loads(text)


def groq_structured(prompt, schema, schema_name, model, max_tokens=2000):
    from groq import Groq
    client = Groq(api_key=GROQ_API_KEY)

    for attempt in range(2):
        try:
            response = client.chat.completions.create(
                model=model,
                messages=[
                    {
                        "role": "system",
                        "content":
                            "Return accurate structured data. "
                            "Never invent facts.",
                    },
                    {"role": "user", "content": prompt},
                ],
                response_format={
                    "type": "json_schema",
                    "json_schema": {
                        "name": schema_name,
                        "strict": True,
                        "schema": schema,
                    },
                },
                reasoning_effort="low",
                include_reasoning=False,
                temperature=0.1,
                max_completion_tokens=max_tokens,
            )
            content = response.choices[0].message.content
            if not content:
                raise ValueError("Groq returned an empty response.")
            return json.loads(content)
        except Exception as error:
            if "429" in str(error) and attempt == 0:
                time.sleep(8)
                continue
            raise


def gemini_structured(prompt, schema, max_tokens=2000):
    from google import genai
    client = genai.Client(api_key=GEMINI_API_KEY)

    schema_text = json.dumps(schema, ensure_ascii=False)
    full_prompt = (
        prompt + "\n\nReturn only JSON matching this schema:\n" + schema_text
    )

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=full_prompt,
        config={
            "response_mime_type": "application/json",
            "temperature": 0.1,
            "max_output_tokens": max_tokens,
        },
    )

    if not response.text:
        raise ValueError("Gemini returned an empty response.")

    return extract_json(response.text)


def structured_call(prompt, schema, schema_name, preferred_model, max_tokens=2000):
    errors = []

    if GROQ_API_KEY:
        models = list(dict.fromkeys(
            [preferred_model, GROQ_JUDGE_MODEL, GROQ_PROFILE_MODEL]
        ))
        for model in models:
            try:
                result = groq_structured(
                    prompt, schema, schema_name, model, max_tokens
                )
                return result, f"Groq: {model}"
            except Exception as error:
                errors.append(f"Groq {model}: {error}")

    if GEMINI_API_KEY:
        try:
            result = gemini_structured(prompt, schema, max_tokens)
            return result, f"Gemini: {GEMINI_MODEL}"
        except Exception as error:
            errors.append(f"Gemini: {error}")

    raise RuntimeError(" | ".join(errors) or "No AI API is configured.")


def conversational_ai(messages):
    errors = []

    if GROQ_API_KEY:
        try:
            from groq import Groq
            response = Groq(api_key=GROQ_API_KEY).chat.completions.create(
                model=GROQ_JUDGE_MODEL,
                messages=messages,
                temperature=0.2,
                max_completion_tokens=1200,
                include_reasoning=False,
            )
            return response.choices[0].message.content.strip()
        except Exception as error:
            errors.append(f"Groq: {error}")

    if GEMINI_API_KEY:
        try:
            from google import genai
            conversation = "\n\n".join(
                f"{message['role'].upper()}:\n{message['content']}"
                for message in messages
            )
            response = genai.Client(api_key=GEMINI_API_KEY).models.generate_content(
                model=GEMINI_MODEL, contents=conversation,
            )
            return response.text.strip()
        except Exception as error:
            errors.append(f"Gemini: {error}")

    return "The AI coach is unavailable.\n\n" + " | ".join(errors)


# =========================================================
# CV READING AND PROFILE EXTRACTION
# =========================================================

def extract_cv_text(file_bytes, filename):
    filename = filename.lower()

    if filename.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        ).strip()

    if filename.endswith(".docx"):
        from docx import Document
        document = Document(io.BytesIO(file_bytes))
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    raise ValueError("Only PDF and DOCX files are supported.")


def understand_cv(cv_text):
    prompt = f"""
Today is {CURRENT_DATE.isoformat()}.

Analyze this CV for scholarship matching.

Rules:
1. Extract only facts supported by the CV.
2. A BSAI or Bachelor of Science in Artificial Intelligence is BS.
3. If the student is completing BS, the normal next degree is MS.
4. Do not confuse project application domains with profession.
5. Short internships do not equal years of full-time work.
6. Extract CGPA/GPA carefully.
7. Never invent nationality, achievements or publications.
8. Use 0 or empty lists for missing values.

CV:

{cv_text[:9000]}
"""

    result, provider = structured_call(
        prompt, CV_SCHEMA, "student_cv_profile", GROQ_PROFILE_MODEL, max_tokens=2600,
    )

    direct_gpa = extract_gpa_from_cv_text(cv_text)
    if direct_gpa > 0:
        result["gpa"] = direct_gpa
        gpa_source = "directly detected from CV text"
    else:
        result["gpa"] = max(0.0, min(4.0, safe_float(result.get("gpa", 0))))
        gpa_source = "AI extraction"

    invalid_names = {"", "none", "unknown", "not specified", "null"}
    if str(result.get("name", "")).strip().lower() in invalid_names:
        result["name"] = fallback_name(cv_text)

    direct_degree = detect_current_degree(cv_text)
    if direct_degree != "unknown":
        result["current_degree"] = direct_degree
    if result["current_degree"] == "BS":
        result["target_degree"] = "MS"

    direct_year = extract_graduation_year(cv_text)
    if direct_year:
        result["expected_graduation_year"] = direct_year

    result["field_of_study"] = normalize_fields(
        result.get("field_of_study", []),
        result.get("major", ""),
        result.get("profession", ""),
    )

    metadata = {
        "provider": provider,
        "gpa_source": gpa_source,
        "direct_gpa": direct_gpa,
    }

    return result, metadata


# =========================================================
# SEMANTIC SEARCH
# =========================================================

@st.cache_resource(show_spinner=False)
def semantic_assets():
    from sentence_transformers import SentenceTransformer
    model = SentenceTransformer(EMBEDDING_MODEL)
    vectors = model.encode(
        DATA["search_text"].tolist(),
        normalize_embeddings=True,
        batch_size=128,
        show_progress_bar=False,
    )
    return model, np.asarray(vectors)


def profile_search_text(profile):
    return (
        f"Student seeking {profile['target_degree']} scholarship. "
        f"Current degree: {profile['current_degree']}. "
        f"Major: {profile['major']}. "
        f"Profession: {profile['profession']}. "
        f"Career goal: {profile['career_goal']}. "
        f"Fields: {', '.join(profile['field_of_study'])}. "
        f"Interests: {', '.join(profile['interests'])}. "
        f"Skills: {', '.join(profile['skills'])}. "
        f"Research experience: {profile['research_experience']}. "
        f"Research topics: {', '.join(profile['research_topics'])}. "
        f"Projects: {', '.join(profile['projects'])}. "
        f"Leadership: {', '.join(profile['leadership'])}. "
        f"Certifications: {', '.join(profile['certifications'])}. "
        f"Achievements: {', '.join(profile['achievements'])}. "
        f"Nationality: {profile['nationality']}. "
        f"Gender: {profile['gender']}. "
        f"GPA: {profile['gpa']}."
    )


def retrieve_candidates(profile, selected_level, include_unclear=False):
    model, vectors = semantic_assets()

    query = model.encode(
        [profile_search_text(profile)], normalize_embeddings=True
    )[0]

    working = DATA.copy()
    working["semantic_score"] = vectors @ query

    working = working[
        working["_levels"].apply(
            lambda levels: level_is_allowed(levels, selected_level, include_unclear)
        )
    ].copy()

    if working.empty:
        return working

    strategic_mask = (
        working["scholarship_name"]
        .str.contains(r"\bKAUST\b|Erasmus Mundus", case=False, regex=True, na=False)
    )

    strategic = working[strategic_mask].sort_values(
        "semantic_score", ascending=False
    ).copy()
    strategic["retrieval_reason"] = "Strategic opportunity"

    curated = working[
        working["is_curated"] & ~strategic_mask
    ].sort_values("semantic_score", ascending=False).head(6).copy()
    curated["retrieval_reason"] = "Curated semantic match"

    raw = working[
        ~working["is_curated"] & ~strategic_mask
    ].sort_values("semantic_score", ascending=False).head(8).copy()
    raw["retrieval_reason"] = "Raw CSV semantic match"

    return (
        pd.concat([strategic, curated, raw], ignore_index=True)
        .drop_duplicates(subset=["scholarship_name"])
        .reset_index(drop=True)
    )


# =========================================================
# ELIGIBILITY JUDGING
# =========================================================

def candidate_payload(row, candidate_id):
    return {
        "candidate_id": candidate_id,
        "name": str(row["scholarship_name"])[:120],
        "description": str(row["description"])[:700],
        "provider": str(row["provider_name"])[:120],
        "location": str(row["location"])[:150],
        "offered_levels": list(row["_levels"]),
        "deadline": str(row["deadline"]),
        "amount": safe_float(row["amount"]),
        "fields": str(row["field_tags"])[:250],
        "verified": bool(row["verified"]),
        "quality_problem": bool(row["data_quality_problem"]),
        "semantic_score": round(safe_float(row["semantic_score"]), 3),
    }


def fallback_judgment(candidate):
    return {
        "candidate_id": candidate["candidate_id"],
        "status": "verify_first",
        "fit_score": max(20, min(60, int(candidate["semantic_score"] * 100))),
        "why":
            "The result is semantically relevant, but the AI eligibility "
            "judgment was unavailable.",
        "missing": ["Current eligibility verification"],
        "actions": ["Run Live Verify before applying"],
        "warnings": ["Fallback classification was used"],
        "priority": "low",
    }


def apply_guardrails(profile, row, verdict):
    verdict.setdefault("missing", [])
    verdict.setdefault("actions", [])
    verdict.setdefault("warnings", [])

    name = str(row["scholarship_name"]).lower()
    description = str(row["description"]).lower()
    combined = name + " " + description + " " + str(row["field_tags"]).lower()

    levels = set(row["_levels"])

    if profile["target_degree"] not in levels and "UNCLEAR" not in levels:
        verdict["status"] = "not_eligible"
        verdict["fit_score"] = min(verdict["fit_score"], 15)
        verdict["priority"] = "low"
        verdict["warnings"].append(
            f"You target {profile['target_degree']}, but this record "
            f"offers {', '.join(levels)}."
        )

    profile_field = (
        profile["major"] + " " + profile["profession"] + " "
        + " ".join(profile["field_of_study"])
    ).lower()

    is_ai_profile = any(
        phrase in profile_field
        for phrase in [
            "artificial intelligence", "machine learning", "computer science",
            "data science", "computer vision", "natural language processing",
            "ai/ml",
        ]
    )

    wrong_field = bool(re.search(
        r"\blita\b|library science|librarian|agriculture|agronomy|"
        r"horticulture|dental student|nursing student|baseball player|"
        r"oratorical contest",
        combined,
    ))

    ai_bridge = bool(re.search(
        r"artificial intelligence|machine learning|computer science|"
        r"data science|computer vision|natural language processing|robotics",
        combined,
    ))

    if is_ai_profile and wrong_field and not ai_bridge:
        verdict["status"] = "wrong_field"
        verdict["fit_score"] = min(verdict["fit_score"], 20)
        verdict["priority"] = "low"
        verdict["warnings"].append(
            "This belongs to a different academic field."
        )

    if is_ai_profile and "helmut schmidt" in name:
        verdict["status"] = "wrong_field"
        verdict["fit_score"] = min(verdict["fit_score"], 25)
        verdict["warnings"].append(
            "This programme focuses on public policy, governance, "
            "law and related subjects."
        )

    if "erasmus mundus" in name:
        if verdict["status"] == "strong":
            verdict["status"] = "possible"
        verdict["fit_score"] = min(verdict["fit_score"], 84)
        verdict["missing"].append(
            "Select a specific AI/CS-related Erasmus programme."
        )

    if "kaust" in name:
        if verdict["status"] == "strong":
            verdict["status"] = "possible"
        verdict["fit_score"] = min(verdict["fit_score"], 88)
        verdict["actions"].append(
            "Verify the current admission round, language requirements "
            "and documents."
        )

    if bool(row["data_quality_problem"]):
        if verdict["status"] in {"strong", "possible"}:
            verdict["status"] = "verify_first"
        verdict["fit_score"] = min(verdict["fit_score"], 64)
        verdict["priority"] = "low"
        verdict["warnings"].append(
            "The CSV deadline or award value is unverified."
        )
        verdict["actions"].append("Run Live Verify before applying.")

    parsed_deadline = pd.to_datetime(row["deadline"], errors="coerce")
    if not pd.isna(parsed_deadline) and parsed_deadline.date() < CURRENT_DATE:
        verdict["status"] = "verify_first"
        verdict["fit_score"] = min(verdict["fit_score"], 60)
        verdict["warnings"].append(
            "The listed deadline has passed. Check whether a new cycle exists."
        )

    verdict["fit_score"] = int(max(0, min(92, verdict["fit_score"])))

    for key in ["missing", "actions", "warnings"]:
        verdict[key] = list(dict.fromkeys(verdict[key]))

    return verdict


def judge_candidates(profile, candidates):
    candidates = candidates.reset_index(drop=True)

    payloads = [
        candidate_payload(row, index)
        for index, (_, row) in enumerate(candidates.iterrows())
    ]

    output = []
    batch_size = 4
    progress = st.progress(0)

    for start in range(0, len(payloads), batch_size):
        batch = payloads[start:start + batch_size]

        prompt = f"""
Today is {CURRENT_DATE.isoformat()}.

Judge these scholarships for the student.

STUDENT:
{json.dumps(profile, ensure_ascii=False, indent=2)}

SCHOLARSHIPS:
{json.dumps(batch, ensure_ascii=False, indent=2)}

Rules:
1. Semantic similarity is not proof of eligibility.
2. Separate relevance, eligibility and readiness.
3. Never invent deadlines, funding or requirements.
4. Wrong-field scholarships must be labelled wrong_field.
5. Missing facts must be listed as missing.
6. Strong means strong relevance, not guaranteed admission.
7. Return one result for every candidate_id.
"""

        returned = {}
        provider = "Fallback rules"

        try:
            response, provider = structured_call(
                prompt, JUDGE_SCHEMA, f"judge_batch_{start}",
                GROQ_JUDGE_MODEL, max_tokens=2200,
            )
            returned = {
                result["candidate_id"]: result for result in response["results"]
            }
        except Exception as error:
            print("Judging error:", error)

        for candidate in batch:
            candidate_id = candidate["candidate_id"]
            verdict = returned.get(candidate_id, fallback_judgment(candidate))
            row = candidates.iloc[candidate_id]
            verdict = apply_guardrails(profile, row, verdict)
            output.append({**row.to_dict(), **verdict, "judge_provider": provider})

        progress.progress(min(1.0, (start + batch_size) / max(len(payloads), 1)))
        time.sleep(1)

    progress.empty()

    status_order = {
        "strong": 0, "possible": 1, "verify_first": 2,
        "wrong_field": 3, "not_eligible": 4, "suspicious": 5,
    }

    output.sort(
        key=lambda result: (
            status_order.get(result["status"], 9), -result["fit_score"],
        )
    )

    return output


# =========================================================
# LIVE SEARCH AND VERIFICATION
# =========================================================

BLOCKED_DOMAINS = [
    "instagram.com", "facebook.com", "reddit.com", "youtube.com",
    "tiktok.com", "linkedin.com", "medium.com",
]


def website_domain(url):
    try:
        return urlparse(url).netloc.lower().replace("www.", "")
    except Exception:
        return ""


def likely_official(url):
    domain = website_domain(url)
    return any(
        marker in domain
        for marker in [".edu", ".ac.", ".gov", ".eu", "kaust.edu.sa"]
    )


def firecrawl_scrape(url):
    if not FIRECRAWL_API_KEY or not url:
        return ""

    response = requests.post(
        "https://api.firecrawl.dev/v2/scrape",
        headers={
            "Authorization": f"Bearer {FIRECRAWL_API_KEY}",
            "Content-Type": "application/json",
        },
        json={
            "url": url,
            "formats": ["markdown"],
            "onlyMainContent": True,
            "timeout": 60000,
        },
        timeout=90,
    )

    if response.status_code >= 400:
        return ""

    result = response.json()
    return result.get("data", result).get("markdown", "")


def live_verify(profile, scholarship):
    if not TAVILY_API_KEY:
        raise RuntimeError("Tavily is not configured.")

    from tavily import TavilyClient

    name = scholarship["scholarship_name"]
    existing_link = str(scholarship.get("link", ""))

    search = TavilyClient(api_key=TAVILY_API_KEY).search(
        query=(
            f'"{name}" official scholarship deadline eligibility '
            f'{CURRENT_DATE.year} {CURRENT_DATE.year + 1}'
        ),
        search_depth="basic",
        max_results=5,
        include_answer=False,
        include_raw_content=False,
        exclude_domains=BLOCKED_DOMAINS,
    )

    results = search.get("results", [])
    best_url = existing_link
    search_content = ""

    if results:
        best = max(
            results,
            key=lambda item: safe_float(item.get("score", 0)) + (
                0.5 if likely_official(item.get("url", "")) else 0
            ),
        )
        if not best_url:
            best_url = best.get("url", "")
        search_content = best.get("content", "")

    page_content = firecrawl_scrape(best_url)

    context = (search_content + "\n\n" + page_content[:7500]).strip()

    if not context:
        raise RuntimeError("No current webpage content was collected.")

    prompt = f"""
Today is {CURRENT_DATE.isoformat()}.

Verify this scholarship using only the provided web content.

STUDENT:
{json.dumps(profile, ensure_ascii=False, indent=2)}

SCHOLARSHIP:
{name}

SOURCE:
{best_url}

CONTENT:
{context[:9000]}

Do not invent missing facts. Use Unknown where necessary.
"""

    result, provider = structured_call(
        prompt, VERIFY_SCHEMA, "scholarship_verification",
        GROQ_JUDGE_MODEL, max_tokens=2200,
    )

    result["source_url"] = best_url
    result["ai_provider"] = provider

    return result


def search_live(profile, level_filter):
    if not TAVILY_API_KEY:
        return []

    from tavily import TavilyClient

    level_terms = {
        "Any": "undergraduate masters PhD",
        "BS": "undergraduate bachelor's",
        "MS": "master's",
        "PhD": "doctoral PhD",
    }

    query = (
        f"official fully funded {level_terms[level_filter]} scholarship "
        f"for {profile['profession'] or profile['major']} "
        f"{profile['nationality']} international student "
        f"{' '.join(profile['field_of_study'])} "
        f"{CURRENT_DATE.year} {CURRENT_DATE.year + 1}"
    )

    response = TavilyClient(api_key=TAVILY_API_KEY).search(
        query=query,
        search_depth="basic",
        max_results=10,
        include_answer=False,
        include_raw_content=False,
        exclude_domains=BLOCKED_DOMAINS,
    )

    output = []
    seen = set()

    for result in response.get("results", []):
        url = result.get("url", "")
        score = safe_float(result.get("score", 0))

        if not url or url in seen or score < 0.35:
            continue

        seen.add(url)

        output.append({
            "title": result.get("title", ""),
            "content": result.get("content", ""),
            "url": url,
            "score": score,
            "official": likely_official(url),
        })

    output.sort(key=lambda item: (item["official"], item["score"]), reverse=True)

    return output[:8]


# =========================================================
# SCHOLARSHIP COACH
# =========================================================

def scholarship_coach_answer(profile, scholarship, verification, question, history):
    context = {
        "student": profile,
        "scholarship": {
            "name": scholarship.get("scholarship_name"),
            "status": scholarship.get("status"),
            "score": scholarship.get("fit_score"),
            "why": scholarship.get("why"),
            "missing": scholarship.get("missing", []),
            "actions": scholarship.get("actions", []),
            "warnings": scholarship.get("warnings", []),
        },
        "verification": verification or "Not live-verified",
    }

    messages = [
        {
            "role": "system",
            "content":
                "You are a scholarship coach. Use only the supplied CV "
                "profile and scholarship context. Never invent "
                "achievements, never guarantee success, and give "
                "specific practical recommendations.",
        },
        {
            "role": "user",
            "content": json.dumps(context, ensure_ascii=False, default=str),
        },
    ]

    messages.extend(history[-8:])
    messages.append({"role": "user", "content": question})

    return conversational_ai(messages)


# =========================================================
# SESSION STATE
# =========================================================

DOCUMENT_OPTIONS = [
    "CV", "Transcript", "Statement of Purpose",
    "Recommendation Letter 1", "Recommendation Letter 2",
    "Recommendation Letter 3", "IELTS", "TOEFL", "Passport",
    "Research Proposal", "Research Paper", "Portfolio",
    "Financial Documents", "Community Service Evidence",
]

DEFAULTS = {
    "student_name": "",
    "gender": "unknown",
    "nationality": "",
    "current_degree": "BS",
    "target_degree": "MS",
    "graduation_year": 2027,
    "gpa": 0.0,
    "major": "",
    "profession": "",
    "career_goal": "",
    "fields": [],
    "interests_text": "",
    "skills_text": "",
    "research_experience": False,
    "research_topics_text": "",
    "projects_text": "",
    "leadership_text": "",
    "certifications_text": "",
    "achievements_text": "",
    "work_years": 0.0,
    "community_hours": 0,
    "ielts": 0.0,
    "toefl": 0,
    "documents": [],
    "level_filter": "MS",
    "include_unclear": False,
    "ai_profile": {},
    "cv_notes": {},
}

for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = (
            value.copy() if isinstance(value, (list, dict)) else value
        )

st.session_state.gpa = max(0.0, min(4.0, safe_float(st.session_state.gpa)))
st.session_state.fields = [
    field for field in st.session_state.fields if field in FIELD_OPTIONS
]

if st.session_state.current_degree not in ["BS", "MS", "PhD"]:
    st.session_state.current_degree = "BS"
if st.session_state.target_degree not in ["BS", "MS", "PhD"]:
    st.session_state.target_degree = "MS"


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.title("👤 Student Profile")

with st.sidebar.expander("🔌 API Status", expanded=True):
    col1, col2 = st.columns(2)
    col1.write("Groq: ✅" if GROQ_API_KEY else "Groq: ❌")
    col2.write("Gemini: ✅" if GEMINI_API_KEY else "Gemini: ❌")
    col1.write("Tavily: ✅" if TAVILY_API_KEY else "Tavily: ❌")
    col2.write("Firecrawl: ✅" if FIRECRAWL_API_KEY else "Firecrawl: ❌")


st.sidebar.subheader("1. Upload CV (optional)")

uploaded_cv = st.sidebar.file_uploader(
    "Upload PDF or DOCX", type=["pdf", "docx"]
)

if uploaded_cv and st.sidebar.button(
    "🧠 Understand My CV", use_container_width=True
):
    try:
        cv_text = extract_cv_text(uploaded_cv.getvalue(), uploaded_cv.name)

        if len(cv_text.strip()) < 100:
            st.sidebar.warning(
                "Very little text was extracted. You can still complete "
                "your profile manually below."
            )
        else:
            with st.spinner("AI is reading your CV..."):
                parsed, metadata = understand_cv(cv_text)

            st.session_state.ai_profile = parsed
            st.session_state.student_name = parsed.get("name") or fallback_name(cv_text)
            st.session_state.gender = parsed.get("gender", "unknown")
            st.session_state.nationality = parsed.get("nationality", "")
            st.session_state.current_degree = parsed.get("current_degree", "BS")
            st.session_state.target_degree = parsed.get("target_degree", "MS")
            st.session_state.level_filter = parsed.get("target_degree", "MS")

            year = safe_int(parsed.get("expected_graduation_year", 0))
            if CURRENT_DATE.year <= year <= 2035:
                st.session_state.graduation_year = year

            st.session_state.gpa = max(
                0.0, min(4.0, safe_float(parsed.get("gpa", 0)))
            )
            st.session_state.major = parsed.get("major", "")
            st.session_state.profession = parsed.get("profession", "")
            st.session_state.career_goal = parsed.get("career_goal", "")

            st.session_state.fields = [
                field for field in normalize_fields(
                    parsed.get("field_of_study", []),
                    parsed.get("major", ""),
                    parsed.get("profession", ""),
                )
                if field in FIELD_OPTIONS
            ]

            st.session_state.interests_text = ", ".join(parsed.get("interests", []))
            st.session_state.skills_text = ", ".join(parsed.get("skills", []))
            st.session_state.research_experience = bool(
                parsed.get("research_experience", False)
            )
            st.session_state.research_topics_text = ", ".join(
                parsed.get("research_topics", [])
            )
            st.session_state.projects_text = "\n".join(parsed.get("projects", []))
            st.session_state.leadership_text = "\n".join(parsed.get("leadership", []))
            st.session_state.certifications_text = "\n".join(
                parsed.get("certifications", [])
            )
            st.session_state.achievements_text = "\n".join(
                parsed.get("achievements", [])
            )
            st.session_state.work_years = safe_float(
                parsed.get("work_experience_years", 0)
            )

            if "CV" not in st.session_state.documents:
                st.session_state.documents.append("CV")

            st.session_state.cv_notes = metadata

            for key in ["match_results", "debug_info", "live_results"]:
                st.session_state.pop(key, None)

            st.sidebar.success(f"CV parsed using {metadata['provider']}")
            st.rerun()

    except Exception as error:
        st.sidebar.warning(
            "Automatic CV parsing is currently unavailable. "
            "You can still complete your profile manually below.\n\n"
            f"Details: {str(error)[:300]}"
        )

if st.session_state.cv_notes:
    direct_gpa = st.session_state.cv_notes.get("direct_gpa", 0)
    if direct_gpa:
        st.sidebar.success(f"✅ GPA directly detected: {direct_gpa:.2f}/4.00")
    else:
        st.sidebar.caption("GPA was taken from AI extraction. Review it below.")


st.sidebar.markdown("---")
st.sidebar.subheader("2. Manual Profile")
st.sidebar.caption(
    "Fill or correct these fields manually — required if CV parsing fails."
)

st.sidebar.text_input("Full name", key="student_name")

st.sidebar.selectbox(
    "Gender",
    options=["female", "male", "non-binary", "unknown"],
    key="gender",
)

st.sidebar.text_input(
    "Nationality", key="nationality", placeholder="Example: Pakistani"
)

st.sidebar.selectbox(
    "Current degree", options=["BS", "MS", "PhD"], key="current_degree"
)

st.sidebar.selectbox(
    "Target degree", options=["BS", "MS", "PhD"], key="target_degree"
)

st.sidebar.number_input(
    "Expected graduation year",
    min_value=2026, max_value=2035, step=1, key="graduation_year",
)

st.sidebar.number_input(
    "GPA on a 4.0 scale",
    min_value=0.0, max_value=4.0, step=0.01,
    format="%.2f", key="gpa",
    help="Enter a GPA between 0.00 and 4.00.",
)

st.sidebar.text_input(
    "Major", key="major", placeholder="Example: Artificial Intelligence"
)

st.sidebar.text_input(
    "Profession", key="profession",
    placeholder="Example: AI/ML Student and Researcher",
)

st.sidebar.text_area(
    "Career goal", key="career_goal",
    placeholder=(
        "Example: Pursue an MS in AI and work on research-driven "
        "machine learning systems."
    ),
    height=100,
)

st.sidebar.markdown("### 🎯 Interests and Skills")

st.sidebar.multiselect(
    "Academic fields",
    options=FIELD_OPTIONS,
    key="fields",
    placeholder="Choose your main fields",
)

st.sidebar.text_area(
    "Interests — comma separated",
    key="interests_text",
    placeholder="AI safety, humanitarian technology, medical imaging",
    height=90,
)

st.sidebar.text_area(
    "Skills — comma separated",
    key="skills_text",
    placeholder="Python, SQL, TensorFlow, Machine Learning, Leadership",
    height=120,
)

st.sidebar.markdown("### 🔬 Research and Projects")

st.sidebar.checkbox("I have research experience", key="research_experience")

st.sidebar.text_area(
    "Research topics — comma separated",
    key="research_topics_text",
    placeholder="Thermal landmine detection, LLM jailbreak classification",
    height=100,
)

st.sidebar.text_area(
    "Projects — one project per line",
    key="projects_text",
    placeholder="Landmine Detection System\nKidney Dataset Annotation",
    height=130,
)

st.sidebar.text_area(
    "Leadership and volunteering — one item per line",
    key="leadership_text",
    placeholder="Stanford Code in Place Section Leader",
    height=110,
)

st.sidebar.text_area(
    "Certifications — one item per line",
    key="certifications_text",
    placeholder="Microsoft AI-102\nStanford Code in Place",
    height=110,
)

st.sidebar.text_area(
    "Awards and achievements — one item per line",
    key="achievements_text",
    placeholder="CGPA 3.8\nSection Leader Certificate",
    height=100,
)

st.sidebar.markdown("### 💼 Experience")

st.sidebar.number_input(
    "Full-time-equivalent work experience",
    min_value=0.0, max_value=30.0, step=0.25, key="work_years",
)

st.sidebar.number_input(
    "Community service hours",
    min_value=0, max_value=10000, step=1, key="community_hours",
)

st.sidebar.markdown("### 📄 Tests and Documents")

st.sidebar.number_input(
    "IELTS — 0 if not taken",
    min_value=0.0, max_value=9.0, step=0.5, key="ielts",
)

st.sidebar.number_input(
    "TOEFL — 0 if not taken",
    min_value=0, max_value=120, step=1, key="toefl",
)

st.sidebar.multiselect(
    "Documents ready", options=DOCUMENT_OPTIONS, key="documents"
)

st.sidebar.markdown("---")
st.sidebar.subheader("3. Scholarship Level")

st.sidebar.selectbox(
    "Show scholarships for",
    options=["Any", "BS", "MS", "PhD"],
    key="level_filter",
)

st.sidebar.checkbox(
    "Include unclear degree-level records", key="include_unclear"
)


# =========================================================
# CURRENT PROFILE
# =========================================================

ai_profile = dict(st.session_state.ai_profile)

profile = {
    "name": st.session_state.student_name or "Student",
    "gender": st.session_state.gender,
    "nationality": st.session_state.nationality,
    "current_degree": st.session_state.current_degree,
    "target_degree": st.session_state.target_degree,
    "expected_graduation_year": int(st.session_state.graduation_year),
    "gpa": float(st.session_state.gpa),
    "major": st.session_state.major,
    "profession": st.session_state.profession,
    "career_goal": st.session_state.career_goal,
    "field_of_study": list(st.session_state.fields),
    "interests": split_comma_values(st.session_state.interests_text),
    "skills": split_comma_values(st.session_state.skills_text),
    "research_experience": bool(st.session_state.research_experience),
    "research_topics": split_comma_values(st.session_state.research_topics_text),
    "projects": split_line_values(st.session_state.projects_text),
    "leadership": split_line_values(st.session_state.leadership_text),
    "certifications": split_line_values(st.session_state.certifications_text),
    "achievements": split_line_values(st.session_state.achievements_text),
    "work_experience_years": float(st.session_state.work_years),
    "community_service_hours": int(st.session_state.community_hours),
    "ielts_score": float(st.session_state.ielts),
    "toefl_score": int(st.session_state.toefl),
    "documents_ready": list(st.session_state.documents),
}


# =========================================================
# HEADER AND TABS
# =========================================================

st.title("🎓 ScholarReady AI")

st.markdown(
    "### CV intelligence → degree filtering → semantic matching "
    "→ eligibility reasoning → live verification"
)

metrics = st.columns(4)
metrics[0].metric("Scholarship Records", f"{len(DATA):,}")
metrics[1].metric("Curated/Anchors", int(DATA["is_curated"].sum()))
metrics[2].metric("Current GPA", f"{profile['gpa']:.2f}/4.00")
metrics[3].metric("Level Filter", st.session_state.level_filter)

st.info(
    "Match scores measure profile relevance, not admission probability "
    "or guaranteed funding."
)

dashboard_tab, match_tab, coach_tab, live_tab, quality_tab = st.tabs([
    "📊 Dashboard", "🎯 Matches", "💬 Scholarship Coach",
    "🌐 Live Search", "📋 Data Quality",
])


# =========================================================
# MATCHES TAB
# =========================================================

with match_tab:
    st.markdown("## 🎯 Intelligent Scholarship Matching")

    with st.expander("Current profile", expanded=True):
        left, right = st.columns(2)

        with left:
            st.write("**Name:**", profile["name"])
            st.write("**Profession:**", profile["profession"])
            st.write("**Major:**", profile["major"])
            st.write(
                "**Degree:**",
                f"{profile['current_degree']} → {profile['target_degree']}",
            )

        with right:
            st.write("**GPA:**", f"{profile['gpa']:.2f}/4.00")
            st.write("**Nationality:**", profile["nationality"])
            st.write("**Fields:**", ", ".join(profile["field_of_study"]) or "Not entered")
            st.write("**Skills:**", ", ".join(profile["skills"]) or "Not entered")
            st.write(
                "**Research:**",
                ", ".join(profile["research_topics"]) or "Not entered",
            )
            st.write(
                "**Projects:**", ", ".join(profile["projects"]) or "Not entered"
            )

    if st.button("🚀 Run Intelligent Match", type="primary", use_container_width=True):
        if not profile["major"]:
            st.error("Enter your major or upload a CV.")
        elif not profile["field_of_study"]:
            st.error("Select at least one field.")
        else:
            with st.spinner("Retrieving relevant scholarships..."):
                candidates = retrieve_candidates(
                    profile, st.session_state.level_filter,
                    st.session_state.include_unclear,
                )

            if candidates.empty:
                st.warning("No scholarships matched this level filter.")
            else:
                names = candidates["scholarship_name"].astype(str).str.lower().tolist()

                st.session_state.debug_info = {
                    "count": len(candidates),
                    "kaust": any("kaust" in name for name in names),
                    "erasmus": any("erasmus mundus" in name for name in names),
                }

                with st.spinner("AI is evaluating eligibility and gaps..."):
                    st.session_state.match_results = judge_candidates(
                        profile, candidates
                    )

                st.session_state.last_profile = dict(profile)

    if "debug_info" in st.session_state:
        debug = st.session_state.debug_info
        cols = st.columns(3)
        cols[0].metric("Candidates Judged", debug["count"])
        cols[1].metric("KAUST Evaluated", "✅ Yes" if debug["kaust"] else "No")
        cols[2].metric("Erasmus Evaluated", "✅ Yes" if debug["erasmus"] else "No")

    if "match_results" in st.session_state:
        results = st.session_state.match_results

        groups = {
            "strong": [], "possible": [], "verify_first": [],
            "wrong_field": [], "not_eligible": [], "suspicious": [],
        }

        for result in results:
            groups[result.get("status", "verify_first")].append(result)

        category_metrics = st.columns(6)
        category_data = [
            ("strong", "🏆 Strong"), ("possible", "🟡 Possible"),
            ("verify_first", "🔎 Verify"), ("wrong_field", "⚠️ Wrong Field"),
            ("not_eligible", "🔴 Not Eligible"), ("suspicious", "🚨 Suspicious"),
        ]

        for column, (key, label) in zip(category_metrics, category_data):
            column.metric(label, len(groups[key]))

        tabs = st.tabs([
            f"🏆 Strong ({len(groups['strong'])})",
            f"🟡 Possible ({len(groups['possible'])})",
            f"🔎 Verify ({len(groups['verify_first'])})",
            f"⚠️ Wrong Field ({len(groups['wrong_field'])})",
            f"🔴 Not Eligible ({len(groups['not_eligible'])})",
            f"🚨 Suspicious ({len(groups['suspicious'])})",
        ])

        def render_results(items, tab):
            with tab:
                if not items:
                    st.info("No results.")
                    return

                for result in items:
                    amount = safe_float(result.get("amount", 0))
                    funding = (
                        f"${amount:,.0f}" if amount > 0
                        else result.get("funding_summary") or "Live verification required"
                    )

                    with st.expander(
                        f"{result['fit_score']}% — {result['scholarship_name']} — {funding}"
                    ):
                        left, right = st.columns([3, 1])

                        with left:
                            st.write("**Provider:**", result.get("provider_name") or "Unknown")
                            st.write("**Level:**", ", ".join(result.get("_levels", [])))
                            st.write("**Deadline:**", result.get("deadline"))
                            st.write("**Location:**", result.get("location"))
                            st.write(result.get("description", "")[:700])

                        with right:
                            st.metric("Relevance", f"{result['fit_score']}%")
                            st.write("**Priority:**", result.get("priority"))
                            st.write("**AI:**", result.get("judge_provider"))

                        st.progress(result["fit_score"] / 100)

                        st.markdown("### Why this result")
                        st.write(result["why"])

                        if result["missing"]:
                            st.markdown("### Missing")
                            for item in result["missing"]:
                                st.write("⚠️", item)

                        if result["actions"]:
                            st.markdown("### Action checklist")
                            for index, action in enumerate(result["actions"]):
                                st.checkbox(
                                    action,
                                    key=f"action_{result['record_id']}_{index}",
                                )

                        if result["warnings"]:
                            st.markdown("### Warnings")
                            for warning in result["warnings"]:
                                st.write("🚩", warning)

                        link = str(result.get("link", ""))
                        if link.startswith("http"):
                            st.link_button("Open Scholarship Page", link)

                        verify_key = f"verification_{result['record_id']}"

                        if st.button("🔎 Live Verify", key=f"button_{verify_key}"):
                            with st.spinner("Searching and verifying..."):
                                try:
                                    st.session_state[verify_key] = live_verify(
                                        profile, result
                                    )
                                except Exception as error:
                                    st.error(str(error))

                        verification = st.session_state.get(verify_key)

                        if verification:
                            st.success("Live verification completed")

                            vcols = st.columns(3)
                            vcols[0].metric(
                                "Verified Fit",
                                f"{verification['student_fit_score']}%",
                            )
                            vcols[1].metric(
                                "Verdict",
                                verification["student_verdict"].replace("_", " ").title(),
                            )
                            vcols[2].metric(
                                "Confidence", verification["confidence"].title()
                            )

                            st.write("**Current deadline:**", verification["current_deadline"])
                            st.write("**Current amount:**", verification["current_amount"])
                            st.write("**Why:**", verification["why_for_student"])

                            if verification["missing_for_student"]:
                                st.write("**Still needed:**")
                                for item in verification["missing_for_student"]:
                                    st.write("•", item)

        render_results(groups["strong"], tabs[0])
        render_results(groups["possible"], tabs[1])
        render_results(groups["verify_first"], tabs[2])
        render_results(groups["wrong_field"], tabs[3])
        render_results(groups["not_eligible"], tabs[4])
        render_results(groups["suspicious"], tabs[5])


# =========================================================
# DASHBOARD TAB
# =========================================================

with dashboard_tab:
    st.markdown("## 📊 Scholarship Dashboard")

    if "match_results" not in st.session_state:
        st.info("Run Intelligent Match to generate analytics.")
    else:
        results = st.session_state.match_results

        chart_data = pd.DataFrame([
            {
                "Scholarship": result["scholarship_name"],
                "Status": result["status"].replace("_", " ").title(),
                "Score": result["fit_score"],
                "Level": ", ".join(result.get("_levels", [])),
                "Source": "Curated" if result.get("is_curated") else "Raw",
            }
            for result in results
        ])

        status_counts = chart_data["Status"].value_counts().reset_index()
        status_counts.columns = ["Status", "Count"]

        left, right = st.columns(2)

        with left:
            figure = px.pie(
                status_counts, names="Status", values="Count",
                hole=0.5, title="Result Categories",
            )
            st.plotly_chart(figure, use_container_width=True)

        with right:
            figure = px.histogram(
                chart_data, x="Score", color="Status", nbins=10,
                title="Relevance Scores",
            )
            st.plotly_chart(figure, use_container_width=True)

        missing_counter = Counter()
        for result in results:
            for item in result.get("missing", []):
                missing_counter[str(item)[:80]] += 1

        if missing_counter:
            gaps = pd.DataFrame(
                missing_counter.most_common(10),
                columns=["Gap", "Scholarships"],
            )
            figure = px.bar(
                gaps.sort_values("Scholarships"),
                x="Scholarships", y="Gap", orientation="h",
                title="Most Common Application Gaps",
            )
            st.plotly_chart(figure, use_container_width=True)

        st.dataframe(
            chart_data.sort_values("Score", ascending=False),
            use_container_width=True, hide_index=True,
        )


# =========================================================
# SCHOLARSHIP COACH TAB
# =========================================================

with coach_tab:
    st.markdown("## 💬 Scholarship Coach")

    if "match_results" not in st.session_state:
        st.info("Run Intelligent Match first.")
    else:
        result_map = {
            result["record_id"]: result
            for result in st.session_state.match_results
        }

        selected_id = st.selectbox(
            "Choose scholarship",
            list(result_map.keys()),
            format_func=lambda record_id: (
                result_map[record_id]["scholarship_name"]
                + " — " + str(result_map[record_id]["fit_score"]) + "%"
            ),
        )

        scholarship = result_map[selected_id]
        verification = st.session_state.get(f"verification_{selected_id}")

        st.write(scholarship.get("why", ""))

        if "coach_threads" not in st.session_state:
            st.session_state.coach_threads = {}

        thread = st.session_state.coach_threads.setdefault(selected_id, [])

        qcols = st.columns(3)
        quick_question = None

        if qcols[0].button("Improve my CV", key=f"cv_{selected_id}"):
            quick_question = "What should I improve in my CV for this scholarship?"

        if qcols[1].button("Why not ready?", key=f"ready_{selected_id}"):
            quick_question = "Why am I not fully ready?"

        if qcols[2].button("30-day plan", key=f"plan_{selected_id}"):
            quick_question = "Give me a realistic 30-day preparation plan."

        for message in thread:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

        typed_question = st.chat_input(
            "Ask about this scholarship or your CV...", key=f"chat_{selected_id}"
        )

        question = typed_question or quick_question

        if question:
            history = list(thread)
            thread.append({"role": "user", "content": question})

            with st.chat_message("user"):
                st.markdown(question)

            with st.chat_message("assistant"):
                with st.spinner("Coach is analyzing..."):
                    answer = scholarship_coach_answer(
                        profile, scholarship, verification, question, history
                    )
                st.markdown(answer)

            thread.append({"role": "assistant", "content": answer})


# =========================================================
# LIVE SEARCH TAB
# =========================================================

with live_tab:
    st.markdown("## 🌐 Live Scholarship Search")

    if st.button("Search Live Scholarships", use_container_width=True):
        with st.spinner("Searching current opportunities..."):
            try:
                st.session_state.live_results = search_live(
                    profile, st.session_state.level_filter
                )
            except Exception as error:
                st.error(str(error))

    for index, result in enumerate(st.session_state.get("live_results", [])):
        badge = "Likely official" if result["official"] else "Third-party source"

        with st.expander(f"{index + 1}. {result['title']} — {badge}"):
            st.write(result["content"])
            st.write("Relevance:", round(result["score"], 3))
            st.link_button("Open Result", result["url"])


# =========================================================
# DATA QUALITY TAB
# =========================================================

with quality_tab:
    st.markdown("## 📋 Dataset Quality")

    qcols = st.columns(4)
    qcols[0].metric("Total Records", f"{len(DATA):,}")
    qcols[1].metric("Curated", int(DATA["is_curated"].sum()))
    qcols[2].metric(
        "Unknown Deadlines",
        int(DATA["deadline"].str.contains("unknown", case=False, na=False).sum()),
    )
    qcols[3].metric("$0/$500 Values", int(DATA["amount"].isin([0, 500]).sum()))

    level_counts = pd.DataFrame({
        "Level": ["BS", "MS", "PhD", "POSTDOC", "UNCLEAR"],
        "Records": [
            int(DATA["_levels"].apply(lambda values: level in values).sum())
            for level in ["BS", "MS", "PhD", "POSTDOC", "UNCLEAR"]
        ],
    })

    figure = px.bar(
        level_counts, x="Level", y="Records",
        title="Scholarships by Degree Level",
    )
    st.plotly_chart(figure, use_container_width=True)

    st.warning(
        "AI-enriched records are not automatically verified. "
        "Always use current official provider information."
    )


st.divider()

st.caption(
    "ScholarReady AI — CV extraction, manual profile input, GPA detection, "
    "semantic matching, AI coaching, Tavily discovery and Firecrawl verification."
)
